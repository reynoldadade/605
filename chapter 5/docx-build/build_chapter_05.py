"""Build chapter_05.docx from chapter-05-draft-sections.md + ch5_pedagogy.py
using the rsc-chapter-writer ChapterBuilder (with the 2026-09-15 Consolas
correction for captions and keywords applied)."""
import re
import sys

sys.path.insert(0, "/home/claude/notes")
sys.path.insert(0, "/home/claude/build")
from chapter_builder import (ChapterBuilder, _set_font, _set_cell_border, _set_cell_margins,
                             FONT_PROSE, FONT_CODE, SZ_CONTENT, SZ_CODE_IN_TEXT, SZ_KEYWORD)
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
import ch5_pedagogy as P
from urllib.parse import quote
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DRAFT = "/home/claude/notes/chapter-05-draft-sections.md"
FIG_DIR = "/home/claude/ch5/figures"
OUT = sys.argv[1] if len(sys.argv) > 1 else "/home/claude/build/chapter_05.docx"

INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")


def runs_of(text, lead_label=False):
    """Split markdown-ish inline text into (text, kind) runs.
    `code` -> code_in_text; **x** -> keyword (or a Lora-bold label when it
    leads a list item); *x* -> italic."""
    out = []
    parts = [p for p in INLINE.split(text) if p]
    for i, part in enumerate(parts):
        if part.startswith("`"):
            out.append((part[1:-1], "code_in_text"))
        elif part.startswith("**"):
            out.append((part[2:-2], "label" if (lead_label and i == 0) else "keyword"))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], "italic"))
        else:
            out.append((part, "normal"))
    return out


STYLE = {
    "normal": (FONT_PROSE, SZ_CONTENT, False, False),
    "keyword": (FONT_CODE, SZ_KEYWORD, True, False),
    "label": (FONT_PROSE, SZ_CONTENT, True, False),
    "italic": (FONT_PROSE, SZ_CONTENT, False, True),
    "code_in_text": (FONT_CODE, SZ_CODE_IN_TEXT, True, False),
}


def write_runs(p, runs):
    for text, kind in runs:
        f, s, b, it = STYLE[kind]
        _set_font(p.add_run(text), f, s, bold=b, italic=it)


class Ch5(ChapterBuilder):
    def para(self, text, lead_label=False):
        p = self.doc.add_paragraph()
        write_runs(p, runs_of(text, lead_label=lead_label))
        return p

    def bullet(self, text):
        text = text.rstrip()
        if text.endswith("."):
            text = text[:-1]
        p = self.doc.add_paragraph(style="List Bullet")
        write_runs(p, runs_of(text, lead_label=True))

    def numbered(self, n, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        _set_font(p.add_run(f"{n}. "), FONT_PROSE, SZ_CONTENT)
        write_runs(p, runs_of(text, lead_label=True))

    def hyperlink(self, p, url, text):
        """A real, clickable hyperlink run (Lora 11, underlined)."""
        r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            fonts.set(qn(a), FONT_PROSE)
        rpr.append(fonts)
        color = OxmlElement("w:color"); color.set(qn("w:val"), "1F4E79"); rpr.append(color)
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(SZ_CONTENT * 2)); rpr.append(sz)
        r.append(rpr)
        t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
        r.append(t)
        h.append(r)
        p._p.append(h)

    def code_ref(self, path):
        """'Full code: chapter 5/<path> (on GitHub)' under a trimmed listing."""
        last = self.doc.paragraphs[-1]
        if not last.text.strip():  # sit directly under the listing
            last._p.getparent().remove(last._p)
        p = self.doc.add_paragraph()
        _set_font(p.add_run("Full code: "), FONT_PROSE, SZ_CONTENT, italic=True)
        _set_font(p.add_run(f"chapter 5/{path}"), FONT_CODE, SZ_CODE_IN_TEXT, bold=True)
        _set_font(p.add_run(" ("), FONT_PROSE, SZ_CONTENT)
        self.hyperlink(p, P.REPO_BLOB + quote(path), "view on GitHub")
        _set_font(p.add_run(")"), FONT_PROSE, SZ_CONTENT)
        self.doc.add_paragraph()

    def add_mcqs(self, questions):
        self.add_heading("Multiple Choice Questions", level=2)
        for i, q in enumerate(questions, 1):
            self.para(f"{i}. {q['question']}")
            for letter, option in zip("ABCD", q["options"]):
                self.para(f"{letter}. {option}")
        self.add_heading("Answer", level=3)
        for i, q in enumerate(questions, 1):
            self.add_body_paragraph(f"{i}  {q['answer']}")

    def add_questions(self, questions):
        self.add_heading("Questions", level=2)
        for i, q in enumerate(questions, 1):
            self.para(f"{i}. {q}")

    def add_key_terms(self, terms):
        """Key terms are keywords being introduced: Consolas 10 bold
        (publisher correction 2026-09-15); definitions in Lora 11."""
        self.add_heading("Key Terms", level=2)
        for term, definition in terms.items():
            p = self.doc.add_paragraph()
            _set_font(p.add_run(term.replace("`", "") + ":"), FONT_CODE, SZ_KEYWORD, bold=True)
            _set_font(p.add_run(" "), FONT_PROSE, SZ_CONTENT)
            write_runs(p, runs_of(definition))

    def callout(self, label, body):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        _set_cell_border(cell)
        _set_cell_margins(cell)
        p = cell.paragraphs[0]
        _set_font(p.add_run(label + " "), FONT_PROSE, SZ_CONTENT, bold=True)
        write_runs(p, runs_of(body))
        self.doc.add_paragraph()

    def md_table(self, header, rows):
        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate([header] + rows):
            for c, val in enumerate(row):
                cell = table.rows[r].cells[c]
                _set_cell_margins(cell, margin_dxa=80)
                p = cell.paragraphs[0]
                if r == 0:
                    _set_font(p.add_run(val), FONT_PROSE, SZ_CONTENT, bold=True)
                else:
                    write_runs(p, runs_of(val))


def body_lines():
    text = open(DRAFT, encoding="utf-8").read()
    start = text.index("## Why Mutation Design Matters")
    end = text.index("\n---\n\n## Still to do")
    return text[start:end].split("\n")


def build():
    ch = Ch5()
    ch.start_chapter(P.CHAPTER_NUMBER, P.CHAPTER_NAME)
    for t in P.INTRODUCTION:
        ch.para(t)
    ch.add_structure(P.STRUCTURE)

    lines = body_lines()
    i = 0
    para_buf = []
    tables_seen = 0

    def flush():
        if para_buf:
            ch.para(" ".join(para_buf))
            para_buf.clear()

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            flush()
            code = []
            i += 1
            while not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            ch.add_code_block("\n".join(code))
        elif line.startswith("## "):
            flush()
            ch.add_heading(line[3:].strip().replace("`", ""), level=1)
        elif line.startswith("### "):
            flush()
            ch.add_heading(line[4:].strip().replace("`", ""), level=2)
        elif line.startswith("#### "):
            flush()
            ch.add_heading(line[5:].strip().replace("`", ""), level=3)
        elif line.startswith("> "):
            flush()
            quote = []
            while i < len(lines) and lines[i].startswith(">"):
                quote.append(lines[i][1:].strip())
                i += 1
            i -= 1
            q = " ".join(quote)
            m = re.match(r"\*\*(Note|Tip)([^*]*?)\.?\*\*\s*(.*)", q)
            assert m, q[:80]
            ch.callout(f"{m.group(1)}{m.group(2).rstrip('.')}:", m.group(3))
        elif line.startswith("- "):
            flush()
            ch.bullet(line[2:])
        elif re.match(r"^\d+\. ", line):
            flush()
            n, rest = line.split(". ", 1)
            ch.numbered(n, rest)
        elif line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            i -= 1
            tables_seen += 1
            ch.md_table(rows[0], rows[1:])
            ch.add_table_caption(5, tables_seen, P.TABLE_5_1_CAPTION)
        elif line.startswith("[[CODE "):
            flush()
            ch.code_ref(line.strip()[7:-2])
        elif line.startswith("[[FIGURE "):
            flush()
            num = line.strip()[9:-2]
            fname, caption = P.FIGURES[num]
            ch.add_figure_image(f"{FIG_DIR}/{fname}", width_inches=6.0)
            ch.add_figure_caption(5, int(num.split(".")[1]), caption)
        elif line.strip() == "":
            flush()
        else:
            para_buf.append(line.strip())
        i += 1
    flush()

    # Conclusion is the last H1 in the draft; Points to Remember follow it.
    ch.add_heading("Points to Remember", level=2)
    for pt in P.POINTS_TO_REMEMBER:
        ch.bullet(pt)

    ch.add_heading("Questions and Exercises", level=1)
    ch.add_heading("Solved Exercises", level=2)
    for n, ex in enumerate(P.SOLVED_EXERCISES, 1):
        ch.para(f"{n}. {ex['problem']}")
        ch.para(f"**Solution:** {ex['solution']}", lead_label=True)
    ch.add_mcqs(P.MCQS)
    ch.add_questions(P.QUESTIONS)
    ch.add_heading("Assignments", level=2)
    for n, a in enumerate(P.ASSIGNMENTS, 1):
        ch.para(f"{n}. {a}")
    ch.add_key_terms(P.KEY_TERMS)
    ch.save(OUT)
    from fonts_finalize import finalize
    finalize(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
